"""Second wave of the 2026-10-05 inbox ingestion (DevLog-021): extract, supplements, Crossref.

Usage: uv run --with python-docx python data/_staging/ingest_2026_10_05/prep2.py [--apply]
"""

import shutil
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(Path(__file__).resolve().parent))
import refresh_metadata as rm  # noqa: E402
from prep import extract  # noqa: E402

APPLY = "--apply" in sys.argv
INBOX = ROOT / "references" / "_inbox"
REFS = ROOT / "references"
NEW = {
    "eltes2020": "10.1038/s41563-020-0725-5", "ogiso2020": "10.1109/jlt.2019.2924671",
    "kohli2023": "10.1109/jlt.2023.3260064", "schwarzenberger2026a": "10.1109/jstqe.2026.3684861",
    "li2024": "10.1109/jlt.2023.3339472", "wang2022": "10.1021/acsphotonics.2c00263",
    "liu2026d": "10.1021/acsphotonics.6c00668", "eltes2023": "10.1364/ofc.2023.th4a.2",
    "schwarzenberger2023a": "10.1364/cleo_si.2023.sth5c.7", "schwarzenberger2023": "10.1049/icp.2023.2357",
    "chen2025": "10.1002/lpor.202500380", "shen2025": "10.1002/lpor.202401092", "zhang2026c": "10.1002/pssr.202500340",
}
SUPP = ["eltes2020", "wang2022", "chen2025", "chelladurai2025"]


def docx_supplement(pid: str) -> None:
    import docx  # python-docx

    src = INBOX / f"{pid}_supplement.docx"
    out = REFS / pid / "supplement"
    out.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(src, out / "source.docx")
    d = docx.Document(str(src))
    paras = [p.text for p in d.paragraphs]
    tables = ["\n".join(" | ".join(c.text for c in row.cells) for row in t.rows) for t in d.tables]
    (out / "text.md").write_text(
        f"---\npaper_id: {pid}\nrole: supplement (docx)\nextraction_method: python-docx paragraphs + tables; "
        "no page markers (cite 'supplement Section/Fig.'); embedded figures not extracted\n---\n\n"
        + "\n".join(paras) + "\n\n## Tables\n\n" + "\n\n".join(tables) + "\n"
    )
    print(f"{pid:22s} supplement (docx)      paras={len(paras)} tables={len(tables)}")


def main() -> None:
    for pid, doi in NEW.items():
        assert (INBOX / f"{pid}.pdf").exists(), pid
        if APPLY:
            extract(pid, INBOX / f"{pid}.pdf", REFS / pid, doi, "primary")
    for pid in SUPP:
        assert (INBOX / f"{pid}_supplement.pdf").exists(), pid
        if APPLY:
            extract(pid, INBOX / f"{pid}_supplement.pdf", REFS / pid / "supplement", "", "supplement")
    if APPLY:
        docx_supplement("shen2025")
        for pid, doi in NEW.items():
            print(pid, "crossref", rm.crossref_by_doi(pid, doi, False))
    else:
        print("dry run ok:", len(NEW), "papers,", len(SUPP), "pdf supplements, 1 docx supplement")


if __name__ == "__main__":
    main()
